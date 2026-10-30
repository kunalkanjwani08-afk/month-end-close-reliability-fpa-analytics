"""Phase 7 - Phase 5 workbook: formula integrity (LibreOffice recalculation of a COPY), chart sources, external
links, and key values versus the independent pandas reference (ref_metrics.json / ref_lines.pkl)."""
import json, re, zipfile
import numpy as np
import pandas as pd
import openpyxl
from p7_common import *

R = Registry()
A = "excel"
REC = WORK / "xlsx_recalc" / "recalculated" / "workbook_copy.xlsx"
R.cmp("X00", A, "Phase 5 workbook SHA-256 unchanged (QA works on a copy)", 1,
      int(sha256(XLSX) == "71e23dca9e0bbce3295223188c296cd3bd5ff51759825487bae049705375f561"), 0)
if not REC.exists():
    R.add("X01", A, "LibreOffice recalculation of the copy", "BLOCKED", note="recalculated copy missing")
    R.save("excel"); raise SystemExit
fw = openpyxl.load_workbook(XLSX)                     # formulas
cw = openpyxl.load_workbook(XLSX, data_only=True)     # cached values written by the generator
rw = openpyxl.load_workbook(REC, data_only=True)      # values recalculated by LibreOffice
nform = nerr = nnocache = ndiff = 0
diffs, errs, refx = [], [], []
for ws in fw.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if not (isinstance(c.value, str) and c.value.startswith("=")):
                continue
            nform += 1
            f = c.value
            if "#REF!" in f:
                refx.append((ws.title, c.coordinate, f[:120]))
            cv = cw[ws.title][c.coordinate].value
            rv = rw[ws.title][c.coordinate].value
            if cv is None:
                nnocache += 1
            if isinstance(rv, str) and rv.startswith("#") or isinstance(cv, str) and cv.startswith("#"):
                nerr += 1; errs.append((ws.title, c.coordinate, cv, rv, f[:100]))
            if isinstance(cv, (int, float)) and isinstance(rv, (int, float)):
                if abs(cv - rv) > 1e-6 * max(1, abs(cv)):
                    ndiff += 1; diffs.append((ws.title, c.coordinate, cv, rv, f[:100]))
            elif cv != rv and not (cv in (None, "") and rv in (None, "")):
                if not (isinstance(cv, bool) or isinstance(rv, bool)) and str(cv) != str(rv):
                    ndiff += 1; diffs.append((ws.title, c.coordinate, cv, rv, f[:100]))
R.add("X01", A, "Formula cells in workbook", "INFO", None, nform)
R.cmp("X02", A, "Formula cells evaluating to an error (#REF!, #DIV/0!, #VALUE!, #N/A, #NAME?) after LibreOffice recalc", 0, nerr, 0,
      note=json.dumps(errs[:10], default=str))
R.cmp("X03", A, "Formulas containing a literal #REF! (broken reference)", 0, len(refx), 0)
R.cmp("X04", A, "Formula cells whose stored value differs from the LibreOffice recalculated value", 0, ndiff, 0,
      tolerance="1e-6 relative", note=json.dumps(diffs[:10], default=str))
R.cmp("X05", A, "Formula cells without a cached value (would show blank in viewers that do not recalc)", 0, nnocache, 0)
z = zipfile.ZipFile(XLSX)
R.cmp("X06", A, "External workbook links", 0, len([n for n in z.namelist() if "externalLink" in n]), 0)
R.cmp("X07", A, "Macros / VBA project", 0, len([n for n in z.namelist() if "vbaProject" in n]), 0)
dn_bad = [n for n in fw.defined_names if "#REF" in str(fw.defined_names[n].attr_text)]
R.cmp("X08", A, "Defined names with broken references", 0, len(dn_bad), 0)
# chart sources
chart_refs = []
for n in z.namelist():
    if n.startswith("xl/charts/chart") and n.endswith(".xml"):
        x = z.read(n).decode("utf-8")
        for f in re.findall(r"<c:f>(.*?)</c:f>", x):
            chart_refs.append((n, f))
badc = []
for n, f in chart_refs:
    import html as _h
    f = _h.unescape(f)
    m = re.match(r"^'?(.*?)'?!\$?([A-Z]+)\$?(\d+)(?::\$?([A-Z]+)\$?(\d+))?$", f)
    if not m:
        badc.append((n, f, "unparsed")); continue
    sh = m.group(1).replace("''", "'")
    if sh not in rw.sheetnames:
        badc.append((n, f, "sheet missing")); continue
    ws = rw[sh]
    rng = f"{m.group(2)}{m.group(3)}" + (f":{m.group(4)}{m.group(5)}" if m.group(4) else "")
    vals = [c.value for row in ws[rng] for c in (row if isinstance(row, tuple) else (row,))] if ":" in rng else [ws[rng].value]
    if all(v is None for v in vals):
        badc.append((n, f, "empty range"))
R.add("X09", A, "Chart series references found", "INFO", None, len(chart_refs), note=f"{len(set(n for n, _ in chart_refs))} charts")
R.cmp("X10", A, "Chart series references pointing to a missing sheet or an empty range", 0, len(badc), 0, note=json.dumps(badc[:10]))

# ---------------------------------------------------------------- values vs independent reference
ref = json.load(open(WORK / "ref_metrics.json"))
L = pd.read_pickle(WORK / "ref_lines.pkl")
P = L[L.posted]
ms = ref["month_section"]
ws = rw["Monthly P&L"]
hdr = None
for i in range(1, 12):
    if ws.cell(i, 1).value == "Month":
        hdr = i; break
mrows = []
for i in range(hdr + 1, hdr + 40):
    d = ws.cell(i, 1).value
    if hasattr(d, "year"):
        ym = f"{d.year}-{d.month:02d}"
        mrows.append((ym, [ws.cell(i, k).value for k in range(3, 10)]))
R.cmp("X20", A, "Monthly P&L: number of month rows", 36, len(mrows), 0)
mx = {"rev": 0, "cogs": 0, "people": 0, "opex": 0, "gp": 0, "proxy": 0, "gm": 0}
bad = 0
for ym, v in mrows:
    rev, cogs, ppl, opx = (ms.get(f"{ym}|{s}", 0.0) for s in ["REVENUE", "COGS", "PEOPLE", "OPEX"])
    exp = {"rev": rev, "cogs": cogs, "gp": rev - cogs, "people": ppl, "opex": opx, "proxy": rev - cogs - ppl - opx}
    got = {"rev": v[0], "cogs": v[1], "gp": v[2], "people": v[4], "opex": v[5], "proxy": v[6]}
    for k in exp:
        d = abs(got[k] - exp[k]); mx[k] = max(mx[k], d); bad += d > TOL_USD
    d = abs(v[3] - (rev - cogs) / rev); mx["gm"] = max(mx["gm"], d); bad += d > TOL_RATIO
R.cmp("X21", A, "Monthly P&L: revenue, COGS, GP S1, GM%, people, opex, result proxy for 36 months vs independent reference", 0, bad, 0,
      population="Posted eligible", tolerance="0.01 USD / 1e-9", note="max abs diff " + json.dumps({k: float(f"{v:.3g}") for k, v in mx.items()}))
# Executive summary headline
es = rw["Executive Summary"]
lab = {es.cell(i, 1).value: es.cell(i, 2).value for i in range(1, es.max_row + 1) if isinstance(es.cell(i, 1).value, str)}
h = ref["headline"]
gpS1 = [g for g in ref["gp"] if g["rowset"] == "ROWS_ALL" and g["treatment"] == "GP_S1_SIGNED_AS_RECORDED" and g["period"] == "ALL YEARS"][0]
for cid, k, exp, tol in [("X22", "Revenue (USD)", h["revenue"], TOL_USD),
                         ("X23", "Gross profit, treatment S1 signed as recorded (USD)", gpS1["gross_profit_usd"], TOL_USD),
                         ("X24", "Gross margin % (S1)", gpS1["gross_margin_pct"], TOL_RATIO),
                         ("X25", "Operating result proxy (USD) - not EBITDA", h["result_proxy"], TOL_USD),
                         ("X26", "Proxy margin %", h["result_proxy"] / h["revenue"], TOL_RATIO),
                         ("X27", "Eligible P&L lines (count)", 58340, 0)]:
    R.cmp(cid, A, f"Executive Summary '{k}' vs independent", exp, lab.get(k), tol, population="Posted eligible")
yr = P[P.sec == "REVENUE"].groupby(P.ym.str[:4]).a.sum()
R.cmp("X28", A, "Executive Summary revenue growth 2024 vs 2023", yr["2024"] / yr["2023"] - 1, lab.get("Revenue growth 2024 vs 2023"), TOL_RATIO)
R.cmp("X29", A, "Executive Summary revenue growth 2025 vs 2024", yr["2025"] / yr["2024"] - 1, lab.get("Revenue growth 2025 vs 2024"), TOL_RATIO)
# Same-sign revenue variance (definition: Posted eligible with budget, actual/budget both non-zero same sign)
B = P[P.b.notna()]
ss = B[(B.a != 0) & (B.b != 0) & (np.sign(B.a) == np.sign(B.b))]
for cid, sec, key in [("X30", "REVENUE", "Revenue variance %, same-sign lines"), ("X31", "COGS", "COGS variance %, same-sign lines"),
                      ("X32", "PEOPLE", "People variance %, same-sign lines"), ("X33", "OPEX", "Opex variance %, same-sign lines")]:
    s = ss[ss.sec == sec]
    R.cmp(cid, A, f"Executive Summary '{key}'", (s.a.sum() - s.b.sum()) / abs(s.b.sum()), lab.get(key), TOL_RATIO,
          definition="(sum actual - sum budget) / |sum budget| on SAME_SIGN_NONZERO lines")
s = ss[ss.sec == "REVENUE"]
R.cmp("X34", A, "Executive Summary 'Revenue variance vs budget, same-sign lines (USD)'", s.a.sum() - s.b.sum(),
      lab.get("Revenue variance vs budget, same-sign lines (USD)"), TOL_USD)
# Gross profit GP1 (24 rows)
gws = rw["Gross Profit"]
TMAP = {"S1 - signed as recorded": "GP_S1_SIGNED_AS_RECORDED", "S2 - discount (4190) at absolute value": "GP_S2_DISCOUNT_ABS",
        "S3 - excluding discounts (4190)": "GP_S3_EXCL_DISCOUNT"}
gp = {(g["rowset"], g["treatment"], str(g["period"])): g for g in ref["gp"]}
n = bad = 0
for i in range(1, 60):
    rs, t, per = gws.cell(i, 1).value, gws.cell(i, 2).value, gws.cell(i, 3).value
    if t in TMAP and per is not None:
        key = ("ROWS_ALL" if rs == "All eligible lines" else "ROWS_ALL_NON_NEGATIVE", TMAP[t], str(per))
        g = gp.get(key)
        if g is None:
            bad += 1; continue
        n += 1
        for col, k in [(4, "rows_in_scope"), (5, "revenue_usd"), (6, "cogs_usd"), (7, "gross_profit_usd")]:
            bad += abs(gws.cell(i, col).value - g[k]) > TOL_USD
        bad += abs(gws.cell(i, 8).value - g["gross_margin_pct"]) > TOL_RATIO
R.cmp("X35", A, "Gross Profit GP1: rows found (2 row sets x 3 treatments x 4 periods)", 24, n, 0)
R.cmp("X36", A, "Gross Profit GP1: lines, revenue, COGS, GP, GM% vs independent S1/S2/S3", 0, bad, 0)
# Entity E1 and division E2
ews = rw["Entity & Division"]
es_ = ref["entity_section"]; ds_ = ref["division_section"]
bad = n = 0
def block(title_prefix):
    for i in range(1, ews.max_row + 1):
        v = ews.cell(i, 1).value
        if isinstance(v, str) and v.startswith(title_prefix):
            j = i + 1
            while ews.cell(j, 3).value != "Revenue (USD)":
                j += 1
            k = j + 1
            out = []
            while ews.cell(k, 1).value not in (None, "Group"):
                out.append(k); k += 1
            return out
for prefix, src in [("E1.", es_), ("E2.", ds_)]:
    for i in block(prefix):
        key = ews.cell(i, 1).value
        n += 1
        rv_, cg, pp, ox = (src.get(f"{key}|{s}", 0.0) for s in ["REVENUE", "COGS", "PEOPLE", "OPEX"])
        for col, v in [(3, rv_), (4, cg), (5, rv_ - cg), (7, pp), (8, ox), (9, rv_ - cg - pp - ox)]:
            bad += abs(ews.cell(i, col).value - v) > TOL_USD
R.cmp("X37", A, "Entity & Division E1/E2: entity and division rows found", 11, n, 0)
R.cmp("X38", A, "Entity & Division E1/E2: revenue, COGS, GP, people, opex, proxy vs independent", 0, bad, 0)
# GL G1
gl = rw["GL & Cost Drivers"]
glr = ref["gl_totals"]
bad = n = 0
for i in range(1, 40):
    v = gl.cell(i, 1).value
    if isinstance(v, (int, float)) and str(int(v)) in glr and gl.cell(i, 5).value is not None and gl.cell(i, 4).value is not None and isinstance(gl.cell(i, 2).value, str) and i < 34:
        n += 1
        bad += abs(gl.cell(i, 5).value - glr[str(int(v))]) > TOL_USD
        bad += gl.cell(i, 4).value != int((P.gl == str(int(v))).sum())
R.cmp("X39", A, "GL & Cost Drivers G1: rows found (24 P&L accounts)", 24, n, 0)
R.cmp("X40", A, "GL & Cost Drivers G1: signed actual and line count per GL vs independent", 0, bad, 0)
# Data Quality Q1 waterfall
dq = rw["Data Quality"]
wf = pd.read_csv(WORK / "exclusions_pl_actual.csv", index_col=0).rows.to_dict()
bad = n = 0
for i in range(1, 20):
    k = dq.cell(i, 1).value
    if k in wf:
        n += 1; bad += dq.cell(i, 2).value != wf[k]
R.cmp("X41", A, "Data Quality Q1 exclusion waterfall rows vs independent re-derivation", 0, bad, 0, note=f"{n} reasons compared")
# Reconciliation sheet self-report vs reality
rq = rw["Reconciliation & QA"]
R.add("X42", A, "Workbook's own reconciliation sheet status (self-reported, not relied on)", "INFO", None, rq["B5"].value)
R.save("excel")
